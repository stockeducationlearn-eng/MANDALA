import hashlib, shutil, re, io
from pathlib import Path
from PIL import Image
import fitz
from docx import Document as DocxDocument
import pytesseract
from database import db, add_chunks
from text_utils import split_paragraphs, _vocabulary

ROOT = Path(__file__).resolve().parent
DOC_DIR = ROOT / "documents"
DOC_DIR.mkdir(exist_ok=True)
SUPPORTED = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def safe_name(name):
    return (re.sub(r'[<>:"/\\\\|?*]', "_", name).strip(" .") or "belge")[:140]

def configure_tesseract(path):
    if path and Path(path).exists():
        pytesseract.pytesseract.tesseract_cmd = path

def ocr_pil(image, lang):
    try:
        return pytesseract.image_to_string(image, lang=lang).strip()
    except Exception as exc:
        return f"[OCR HATASI: {exc}]"

def extract_file(path, tesseract_path="", ocr_lang="tur+eng"):
    path = Path(path)
    ext = path.suffix.lower()
    chunks = []
    configure_tesseract(tesseract_path)
    if ext == ".pdf":
        with fitz.open(path) as doc:
            pages = []
            for i, page in enumerate(doc, 1):
                text = page.get_text("text").strip()
                if len(text) < 20 and tesseract_path and Path(tesseract_path).exists():
                    try:
                        pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                        text = ocr_pil(Image.open(io.BytesIO(pix.tobytes("png"))), ocr_lang)
                    except Exception as exc:
                        text = text or f"[Sayfa OCR hatası: {exc}]"
                if text:
                    pages.append((i, text))
            # Kelime onarımı belgenin tümüne bakılarak yapılır; tek sayfa
            # sözlüğü seyrek kalır ve satır sonu bölünmelerini kaçırır.
            vocab = _vocabulary("\n".join(t for _, t in pages))
            for i, text in pages:
                paras = split_paragraphs(text, vocabulary=vocab)
                for j, para in enumerate(paras, 1):
                    loc = f"PDF sayfa {i}" if len(paras) == 1 else f"PDF sayfa {i}, paragraf {j}"
                    chunks.append((loc, para))
    elif ext == ".docx":
        doc = DocxDocument(path)
        for i, para in enumerate(doc.paragraphs, 1):
            if para.text.strip():
                chunks.append((f"Paragraf {i}", para.text.strip()))
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if txt:
                    chunks.append((f"Tablo {ti}, satır {ri}", txt))
    elif ext == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        for i, para in enumerate(split_paragraphs(text), 1):
            chunks.append((f"Paragraf {i}", para))
    elif ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        if tesseract_path and Path(tesseract_path).exists():
            chunks.append(("Görsel", ocr_pil(Image.open(path), ocr_lang)))
        else:
            chunks.append(("Görsel", "[OCR kullanılmadı. Ayarlardan Tesseract OCR yolunu belirtin.]"))
    return chunks

def add_file(source_path, tesseract_path="", ocr_lang="tur+eng"):
    source = Path(source_path)
    if source.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Desteklenmeyen dosya türü: {source.suffix}")
    digest = sha256_file(source)
    with db() as con:
        old = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()
        if old:
            return old["id"], "Bu dosya daha önce eklenmiş; mevcut kayıt kullanıldı."
    dest = DOC_DIR / f"{digest[:12]}_{safe_name(source.stem)}{source.suffix.lower()}"
    if not dest.exists():
        shutil.copy2(source, dest)
    extracted = extract_file(dest, tesseract_path, ocr_lang)
    char_count = sum(len(t) for _, t in extracted)
    with db() as con:
        cur = con.execute("""INSERT INTO documents(name, stored_path, file_type, sha256, added_at, char_count, status)
            VALUES (?, ?, ?, ?, datetime('now','localtime'), ?, ?)""",
            (source.name, str(dest), source.suffix.lower(), digest, char_count,
             "Metin çıkarıldı" if char_count else "Metin bulunamadı"))
        doc_id = cur.lastrowid
        add_chunks(con, doc_id, extracted)
    return doc_id, f"Arşive eklendi; {len(extracted)} metin bölümü çıkarıldı."


def _store_text(name, file_type, source_text, origin, digest):
    """Metni arşive yazar ve indeksler.

    `stored_path` her zaman yerel `.txt` kopyasını tutar; `source_url` ise
    metnin alındığı adresi (varsa) saklar. Böylece kaynak silinebilir ve
    `Kaynak dosyayı aç` orijinal adresi açabilir.
    """
    dest = DOC_DIR / f"{digest[:12]}_{safe_name(name)}.txt"
    if not dest.exists():
        dest.write_text(source_text, encoding="utf-8")
    # Satır satır kaydetmek yerine paragraf birleştirilir: satır başına parça,
    # bir cümleyi 10+ ayrı kaynağa böler ve aramayı zayıflatır.
    paragraphs = split_paragraphs(source_text)
    if not paragraphs:
        paragraphs = [source_text.strip() or "(boş)"]
    lines = [(f"Paragraf {i}", para) for i, para in enumerate(paragraphs, 1)]
    with db() as con:
        cur = con.execute("""INSERT INTO documents(name, stored_path, file_type, sha256, added_at, char_count, status, source_url)
            VALUES (?, ?, ?, ?, datetime('now','localtime'), ?, ?, ?)""",
            (name, str(dest), file_type, digest, len(source_text),
             "Metin çıkarıldı" if source_text.strip() else "Metin bulunamadı", origin or ""))
        doc_id = cur.lastrowid
        add_chunks(con, doc_id, lines)
    return doc_id, f"Arşive eklendi; {len(lines)} metin bölümü çıkarıldı."


def add_url(url, timeout=None):
    """Bir web adresinin metnini indirip arşive kaynak olarak ekler."""
    from text_sources import fetch_url_text
    text = fetch_url_text(url) if timeout is None else fetch_url_text(url, timeout=timeout)
    if not text.strip():
        raise ValueError("Sayfada okunabilir metin bulunamadı.")
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    with db() as con:
        old = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()
        if old:
            return old["id"], "Bu sayfa daha önce eklenmiş; mevcut kayıt kullanıldı."
    clean = url.strip()
    name = re.sub(r"^https?://", "", clean).split("?")[0][:80] or "web_sayfasi"
    return _store_text(name, "web", text, clean, digest)


def add_text(text, name="Serbest metin"):
    """Kullanıcının yapıştırdığı/ekrandan gelen metni arşive ekler."""
    if not text or not text.strip():
        raise ValueError("Eklenecek metin boş.")
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    with db() as con:
        old = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()
        if old:
            return old["id"], "Bu metin daha önce eklenmiş; mevcut kayıt kullanıldı."
    return _store_text(name, "text", text, "", digest)
