from pathlib import Path

import pytest

import ingest


def test_extract_txt_splits_lines(ingest_dir, tmp_path):
    src = tmp_path / "belge.txt"
    src.write_text("Birinci satır\n\nİkinci satır\n", encoding="utf-8")
    chunks = ingest_dir.extract_file(src)
    assert [loc for loc, _ in chunks] == ["Satır 1", "Satır 3"]
    assert chunks[0][1] == "Birinci satır"


def test_extract_docx_paragraphs_and_tables(ingest_dir, tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Tazminat hakkı")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Koşul"
    table.rows[0].cells[1].text = "Süre"
    src = tmp_path / "belge.docx"
    doc.save(src)
    chunks = ingest_dir.extract_file(src)
    locations = [loc for loc, _ in chunks]
    assert "Paragraf 1" in locations
    assert any(loc.startswith("Tablo 1") for loc in locations)
    assert any("Koşul | Süre" in text for _, text in chunks)


def test_extract_pdf_pages(ingest_dir, tmp_path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Kidem tazminati sartlari burada aciklanmistir.")
    src = tmp_path / "belge.pdf"
    doc.save(src)
    doc.close()
    chunks = ingest_dir.extract_file(src)
    assert chunks[0][0] == "PDF sayfa 1"
    assert "tazminati" in chunks[0][1]


def test_add_file_indexes_chunks_and_dedupes(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.txt"
    src.write_text("İşçi kıdem tazminatı alır.\n", encoding="utf-8")
    doc_id, msg = ingest_dir.add_file(src)
    assert "eklendi" in msg
    with temp_db.db() as con:
        chunks = con.execute("SELECT COUNT(*) FROM chunks WHERE document_id=?", (doc_id,)).fetchone()[0]
        fts = con.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0]
    assert chunks == 1 and fts == 1

    again_id, again_msg = ingest_dir.add_file(src)
    assert again_id == doc_id
    assert "daha önce" in again_msg
    with temp_db.db() as con:
        total = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    assert total == 1


def test_add_file_rejects_unsupported(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.csv"
    src.write_text("a;b", encoding="utf-8")
    with pytest.raises(ValueError):
        ingest_dir.add_file(src)


def test_delete_document_removes_fts_rows(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.txt"
    src.write_text("Sağlık hizmeti ödemesi.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src)
    with temp_db.db() as con:
        stored = temp_db.delete_document(doc_id)
        fts = con.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0]
    assert stored
    assert fts == 0
