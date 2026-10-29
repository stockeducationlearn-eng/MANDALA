from pathlib import Path

import pytest



def test_extract_txt_splits_paragraphs(ingest_dir, tmp_path):
    src = tmp_path / "belge.txt"
    src.write_text("Birinci paragraf\n\nİkinci paragraf\n", encoding="utf-8")
    chunks = ingest_dir.extract_file(src)
    assert [loc for loc, _ in chunks] == ["Paragraf 1", "Paragraf 2"]
    assert chunks[0][1] == "Birinci paragraf"


def test_extract_txt_joins_wrapped_lines(ingest_dir, tmp_path):
    """PDF/OCR çıktısı gibi sarılmış satırlar tek paragrafta birleşmeli."""
    src = tmp_path / "sarilmis.txt"
    src.write_text("Manevi tazminat iki hâlde öde\nnir. Bu ödenir.\n", encoding="utf-8")
    chunks = ingest_dir.extract_file(src)
    assert len(chunks) == 1
    assert chunks[0][1] == "Manevi tazminat iki hâlde ödenir. Bu ödenir."


def test_extract_txt_unknown_broken_word_stays_spaced(ingest_dir, tmp_path):
    """Belgede bütün hâli hiç geçmeyen bölünme onarılamaz; boşlukla birleşir.

    Yanlış birleştirme (iki kelimeyi birleştirmek) daha büyük hata olduğu için
    emin olunmayan durumda boşluk korunur.
    """
    src = tmp_path / "bilinmeyen.txt"
    src.write_text("Burada zzqq\nyyww yazıyor.\n", encoding="utf-8")
    chunks = ingest_dir.extract_file(src)
    assert chunks[0][1] == "Burada zzqq yyww yazıyor."


def test_extract_docx_paragraphs_and_tables(ingest_dir, tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Tazminat hakkı")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Koşul"
    table.rows[0].cells[1].text = "Süre"
    src = tmp_path / "belge.docx"
    doc.save(src)
    chunks = ingest_dir.extract_file(src)
    locations = [loc for loc, _ in chunks]
    assert "Paragraf 1" in locations
    assert any(loc.startswith("Tablo 1") for loc in locations)
    assert any("Koşul | Süre" in text for _, text in chunks)


def test_extract_pdf_pages(ingest_dir, tmp_path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Kidem tazminati sartlari burada aciklanmistir.")
    src = tmp_path / "belge.pdf"
    doc.save(src)
    doc.close()
    chunks = ingest_dir.extract_file(src)
    assert chunks[0][0] == "PDF sayfa 1"
    assert "tazminati" in chunks[0][1]


def test_add_file_indexes_chunks_and_dedupes(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.txt"
    src.write_text("İşçi kıdem tazminatı alır.\n", encoding="utf-8")
    doc_id, msg = ingest_dir.add_file(src)
    assert "eklendi" in msg
    with temp_db.db() as con:
        chunks = con.execute("SELECT COUNT(*) FROM chunks WHERE document_id=?", (doc_id,)).fetchone()[0]
        fts = con.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0]
    assert chunks == 1 and fts == 1

    again_id, again_msg = ingest_dir.add_file(src)
    assert again_id == doc_id
    assert "daha önce" in again_msg
    with temp_db.db() as con:
        total = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    assert total == 1


def test_add_file_rejects_unsupported(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.csv"
    src.write_text("a;b", encoding="utf-8")
    with pytest.raises(ValueError):
        ingest_dir.add_file(src)


def test_delete_document_removes_fts_rows(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "kaynak.txt"
    src.write_text("Sağlık hizmeti ödemesi.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src)
    with temp_db.db() as con:
        stored = temp_db.delete_document(doc_id)
        fts = con.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0]
    assert stored
    assert fts == 0


def test_add_url_archives_and_indexes(temp_db, ingest_dir, web_server):
    url = web_server + "/sayfa"
    doc_id, msg = ingest_dir.add_url(url)
    assert "eklendi" in msg
    with temp_db.db() as con:
        row = con.execute("SELECT name, stored_path, file_type, source_url FROM documents WHERE id=?", (doc_id,)).fetchone()
    assert row["file_type"] == "web"
    # stored_path yerel kopya (silinebilir), source_url ise orijinal adres.
    assert row["source_url"] == url
    assert row["stored_path"].endswith(".txt")
    assert Path(row["stored_path"]).exists()
    from search_engine import search
    _, results = search("tazminat")
    assert any("tazminatı" in r["text"] for r in results)
    assert results[0]["source_url"] == url


def test_add_url_dedupes_same_page(temp_db, ingest_dir, web_server):
    first, _ = ingest_dir.add_url(web_server + "/sayfa")
    second, msg = ingest_dir.add_url(web_server + "/sayfa")
    assert first == second
    assert "daha önce" in msg


def test_add_text_archives_user_input(temp_db, ingest_dir):
    doc_id, msg = ingest_dir.add_text("İşçi kıdem tazminatı alır.", name="Not")
    assert "eklendi" in msg
    with temp_db.db() as con:
        row = con.execute("SELECT file_type FROM documents WHERE id=?", (doc_id,)).fetchone()
    assert row["file_type"] == "text"


def test_add_text_rejects_empty(temp_db, ingest_dir):
    with pytest.raises(ValueError):
        ingest_dir.add_text("   ")


def test_add_file_new_source_joins_study_list(temp_db, ingest_dir, tmp_path):
    """Yeni eklenen dosya çalışma listesine girer, arşive değil."""
    src = tmp_path / "yeni.txt"
    src.write_text("Yeni çalışma kaynağı.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src)
    with temp_db.db() as con:
        row = con.execute("SELECT in_study, in_archive FROM documents WHERE id=?", (doc_id,)).fetchone()
    assert (row["in_study"], row["in_archive"]) == (1, 0)


def test_add_file_to_archive_flag(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "arsiv.txt"
    src.write_text("Arşive giden kaynak.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src, to_archive=True)
    with temp_db.db() as con:
        row = con.execute("SELECT in_study, in_archive FROM documents WHERE id=?", (doc_id,)).fetchone()
    assert (row["in_study"], row["in_archive"]) == (0, 1)


def test_add_file_same_name_different_content_raises_conflict(temp_db, ingest_dir, tmp_path):
    """Aynı isimde farklı içerik: işlem engellenir ve yeni ad önerilir."""
    a = tmp_path / "rapor.txt"
    a.write_text("Birinci içerik.\n", encoding="utf-8")
    ingest_dir.add_file(a)
    b = tmp_path / "rapor.txt"
    b.write_text("Tamamen başka bir içerik.\n", encoding="utf-8")
    with pytest.raises(ingest_dir.NameConflictError) as info:
        ingest_dir.add_file(b)
    assert info.value.conflict["name"] == "rapor.txt"
    assert info.value.suggestion == "rapor (2).txt"


def test_add_file_conflict_can_be_resolved_with_new_name(temp_db, ingest_dir, tmp_path):
    a = tmp_path / "rapor.txt"
    a.write_text("Birinci içerik.\n", encoding="utf-8")
    ingest_dir.add_file(a)
    b = tmp_path / "rapor.txt"
    b.write_text("İkinci, farklı içerik.\n", encoding="utf-8")
    doc_id, msg = ingest_dir.add_file(b, display_name="rapor (2).txt")
    assert "eklendi" in msg
    with temp_db.db() as con:
        names = [r["name"] for r in con.execute("SELECT name FROM documents ORDER BY id")]
    assert names == ["rapor.txt", "rapor (2).txt"]


def test_add_file_conflict_callback_can_cancel(temp_db, ingest_dir, tmp_path):
    a = tmp_path / "rapor.txt"
    a.write_text("Birinci içerik.\n", encoding="utf-8")
    ingest_dir.add_file(a)
    b = tmp_path / "rapor.txt"
    b.write_text("Başka içerik.\n", encoding="utf-8")
    doc_id, msg = ingest_dir.add_file(b, on_conflict=lambda _c: None)
    assert doc_id is None
    assert "iptal" in msg
    with temp_db.db() as con:
        total = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    assert total == 1


def test_add_file_identical_content_is_reused_not_duplicated(temp_db, ingest_dir, tmp_path):
    """Aynı içerik, farklı adla gelse de ikinci kopya oluşmaz."""
    a = tmp_path / "rapor.txt"
    a.write_text("Aynı içerik.\n", encoding="utf-8")
    first, _ = ingest_dir.add_file(a)
    b = tmp_path / "rapor_kopya.txt"
    b.write_text("Aynı içerik.\n", encoding="utf-8")
    second, msg = ingest_dir.add_file(b)
    assert first == second
    assert "daha önce" in msg
    with temp_db.db() as con:
        total = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    assert total == 1


def test_rename_document_updates_name(temp_db, ingest_dir, tmp_path):
    src = tmp_path / "eski.txt"
    src.write_text("Metin.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src)
    temp_db.rename_document(doc_id, "yeni ad.txt")
    with temp_db.db() as con:
        name = con.execute("SELECT name FROM documents WHERE id=?", (doc_id,)).fetchone()["name"]
    assert name == "yeni ad.txt"


def test_study_and_archive_lists_are_disjoint(temp_db, ingest_dir, tmp_path):
    a = tmp_path / "a.txt"
    a.write_text("A içeriği.\n", encoding="utf-8")
    b = tmp_path / "b.txt"
    b.write_text("B içeriği.\n", encoding="utf-8")
    id_a, _ = ingest_dir.add_file(a)
    id_b, _ = ingest_dir.add_file(b, to_archive=True)
    assert [d["id"] for d in temp_db.list_study_documents()] == [id_a]
    assert [d["id"] for d in temp_db.list_archive_documents()] == [id_b]


def test_migrated_rows_do_not_appear_in_both_lists(temp_db, ingest_dir, tmp_path):
    """Eski kayıt (in_study=1, in_archive=1) yalnızca çalışma listesinde görünür."""
    src = tmp_path / "eski.txt"
    src.write_text("Eski kayıt.\n", encoding="utf-8")
    doc_id, _ = ingest_dir.add_file(src)
    with temp_db.db() as con:
        con.execute("UPDATE documents SET in_archive=1 WHERE id=?", (doc_id,))
    assert [d["id"] for d in temp_db.list_study_documents()] == [doc_id]
    assert temp_db.list_archive_documents() == []
