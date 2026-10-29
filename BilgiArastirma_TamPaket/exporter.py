import csv, zipfile
from pathlib import Path
from datetime import datetime
from database import DB_PATH

ROOT = Path(__file__).resolve().parent
EXPORT_DIR = ROOT / "exports"
EXPORT_DIR.mkdir(exist_ok=True)

def _safe_name(text):
    return "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in (text or "arama")).strip()[:60] or "arama"

# exports klasöründeki dosya uzantıları ve okunabilir tür adları.
REPORT_KINDS = {
    ".pdf": "PDF raporu", ".docx": "Word raporu", ".txt": "Metin / konu notu",
    ".csv": "CSV tablo", ".zip": "Yedek",
}


def report_kind(path):
    """Rapor dosyasının tür etiketini döndürür."""
    return REPORT_KINDS.get(Path(path).suffix.lower(), "Dosya")


def list_reports():
    """exports klasöründeki raporları en yeni önce listeler.

    Kullanıcı oluşturduğu raporları uygulama içinden bulabilsin diye
    dosya sistemi tek kaynak kabul edilir; ayrıca bir kayıt tutulmaz.
    """
    if not EXPORT_DIR.exists():
        return []
    reports = []
    for p in EXPORT_DIR.rglob("*"):
        if not p.is_file() or p.name.startswith("."):
            continue
        stat = p.stat()
        reports.append({
            "path": p, "name": p.name, "kind": report_kind(p),
            "size": stat.st_size, "modified": datetime.fromtimestamp(stat.st_mtime),
        })
    reports.sort(key=lambda r: r["modified"], reverse=True)
    return reports


def human_size(size):
    """Dosya boyutunu okunabilir biçime çevirir (ör. 1,2 MB)."""
    size = float(size or 0)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            value = f"{size:.0f}" if unit == "B" else f"{size:.1f}"
            return value.replace(".", ",") + " " + unit
        size /= 1024

def _report_text(results, query="", mode="", ai_answer=""):
    lines = ["BİLGİ ARAŞTIRMA — ARAMA RAPORU", "=" * 60,
             f"Arama / girdi: {query}", f"Araştırma türü: {mode}",
             f"Oluşturulma: {datetime.now():%Y-%m-%d %H:%M:%S}",
             f"Bulunan kaynak bölümü: {len(results)}", ""]
    if ai_answer.strip():
        lines += ["OLUŞTURULAN KONU NOTU / YAPAY ZEKÂ AÇIKLAMASI", "-" * 60, ai_answer.strip(), "", "KAYNAK SONUÇLARI", "-" * 60]
    for i, r in enumerate(results, 1):
        lines += [f"{i}. {r.get('document_name', 'Kaynak')}",
                  f"Konum: {r.get('location', 'Belirtilmemiş')} | Puan: {round(float(r.get('score', 0)), 2)}",
                  r.get('text', ''), ""]
        if r.get('note'):
            lines += [f"Kullanıcı notu: {r['note']}", ""]
    lines += ["", "Not: Kaynak sayfa/konumları arşivdeki metin çıkarımına göre gösterilir. Yapay zekâ açıklamasını özgün kaynaklarla karşılaştırın."]
    return "\n".join(lines)

def export_results(results, fmt="txt", query="", mode="", ai_answer=""):
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = EXPORT_DIR / f"Arama_Raporu_{_safe_name(query)}_{stamp}.{fmt}"
    body = _report_text(results, query, mode, ai_answer)
    if fmt == "txt":
        path.write_text(body, encoding="utf-8-sig")
    elif fmt == "docx":
        from docx import Document
        doc = Document()
        for line in body.splitlines():
            if line.startswith("BİLGİ ARAŞTIRMA"):
                doc.add_heading(line, 0)
            elif line and set(line) == {"="}:
                continue
            elif line and set(line) == {"-"}:
                continue
            elif line.startswith(("OLUŞTURULAN KONU NOTU", "KAYNAK SONUÇLARI")):
                doc.add_heading(line, level=1)
            elif line[:1].isdigit() and ". " in line[:5]:
                doc.add_heading(line, level=2)
            else:
                doc.add_paragraph(line)
        doc.save(path)
    elif fmt == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib import colors
        # Önce repoyla gelen Unicode font, sonra işletim sistemindeki yaygın fontlar.
        font_paths = [
            ROOT / "fonts" / "DejaVuSans.ttf",
            Path(r"C:\Windows\Fonts\arial.ttf"),
            Path(r"C:\Windows\Fonts\calibri.ttf"),
            Path(r"C:\Windows\Fonts\segoeui.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/Library/Fonts/Arial Unicode.ttf"),
        ]
        font_name = "Helvetica"
        for fp in font_paths:
            if fp.exists():
                try:
                    pdfmetrics.registerFont(TTFont("LocalUnicode", str(fp)))
                    font_name = "LocalUnicode"
                    break
                except Exception:
                    pass
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="TRBody", parent=styles["BodyText"], fontName=font_name, fontSize=9, leading=12, spaceAfter=6))
        styles.add(ParagraphStyle(name="TRTitle", parent=styles["Title"], fontName=font_name, fontSize=16, leading=20, spaceAfter=12))
        styles.add(ParagraphStyle(name="TRHead", parent=styles["Heading2"], fontName=font_name, fontSize=11, leading=14, textColor=colors.HexColor("#243b53"), spaceBefore=8, spaceAfter=5))
        from xml.sax.saxutils import escape
        story = []
        for i, line in enumerate(body.splitlines()):
            if i == 0:
                story.append(Paragraph(escape(line), styles["TRTitle"]))
            elif line.startswith(("OLUŞTURULAN KONU NOTU", "KAYNAK SONUÇLARI")):
                story.append(Paragraph(escape(line), styles["TRHead"]))
            elif line and set(line) <= {"=", "-"}:
                continue
            elif line[:1].isdigit() and ". " in line[:5]:
                story.append(Paragraph(escape(line), styles["TRHead"]))
            elif line:
                story.append(Paragraph(escape(line), styles["TRBody"]))
            else:
                story.append(Spacer(1, 4))
        SimpleDocTemplate(str(path), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42, title=f"Arama Raporu - {query}").build(story)
    else:
        raise ValueError("Desteklenmeyen çıktı türü: " + fmt)
    return path

def export_csv(results):
    path = EXPORT_DIR / f"arama_sonuclari_{datetime.now():%Y%m%d_%H%M%S}.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Kaynak", "Konum", "Puan", "Metin", "Not"])
        for r in results:
            writer.writerow([r["document_name"], r["location"], round(r["score"], 2), r["text"], r.get("note", "")])
    return path

def backup_project():
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = EXPORT_DIR / f"yedek_{stamp}.zip"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as z:
        if DB_PATH.exists(): z.write(DB_PATH, "data/pdf_arastirma.db")
        settings = ROOT / "data" / "settings.json"
        if settings.exists(): z.write(settings, "data/settings.json")
        docs = ROOT / "documents"
        if docs.exists():
            for p in docs.rglob("*"):
                if p.is_file(): z.write(p, str(Path("documents") / p.relative_to(docs)))
    return target
